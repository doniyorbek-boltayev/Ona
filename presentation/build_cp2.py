"""Builds the three Checkpoint 2 decks (texnik, biznes, soha mentor), one slide per judging criterion.

    .venv/bin/python presentation/build_cp2.py

Market figures and the business model follow the team's Checkpoint 1 deck (sources named on the slides).
Numbers marked "gipoteza" are assumptions to be validated in a pilot, not measurements.
"""
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PRODUCT, TEAM, DATE = "Nigoh AI", "Anicenna jamoasi", "19-sentabr 2026"
REPO = "github.com/doniyorbek-boltayev/Ona"

TEAL, TEAL_DARK, DARK = "0E6B78", "0A4F59", "15262A"
INK, MUTED, SOFT, LIGHT, WHITE = "1E2A2E", "5B6B71", "DCF0F1", "F3F7F7", "FFFFFF"
CORAL, GREEN, RED, PALE = "E8846B", "2B9A66", "CF3A30", "9DB7BB"
FONT = "Arial"


def uz(text):
    return re.sub(r"([oOgG])'", "\\1\u2018", text).replace("'", "\u2019")


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, mentor):
        self.mentor = mentor
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(self, bg=WHITE, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        if notes:
            s.notes_slide.notes_text_frame.text = uz(notes)
        return s

    def save(self, name):
        path = os.path.join(HERE, name)
        self.prs.save(path)
        print(path)


def box(s, x, y, w, h, fill=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.08):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(content if isinstance(content, list) else [content]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        if i:
            p.space_before = Pt(size * 0.4)
        for run_text, opts in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = uz(run_text)
            r.font.name, r.font.size = FONT, Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", bold)
            r.font.color.rgb = rgb(opts.get("color", color))
    return tb


def badge(s, x, y, label, d=0.62, fill=TEAL, color=WHITE, size=20):
    box(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    text(s, x, y, d, d, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def picture(s, name, x, y, w):
    box(s, x - 0.08, y - 0.08, w + 0.16, w * 812 / 1491 + 0.16, fill="E3E8EA", radius=0.03)
    s.shapes.add_picture(os.path.join(HERE, "img", name), Inches(x), Inches(y), width=Inches(w))


def bullets(s, x, y, w, items, size=13.5, gap=0.44, color=INK, dot=CORAL):
    for j, item in enumerate(items):
        box(s, x, y + 0.09 + j * gap, 0.13, 0.13, fill=dot, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.3, y + j * gap, w - 0.3, gap, item, size=size, color=color)


# ---- slide templates ---------------------------------------------------------------

def title_slide(d, tagline, image="stage.jpg"):
    s = d.slide(DARK, notes="Salom. Biz %s: homilador ayol har kuni 2 daqiqalik so'rovnoma to'ldiradi, xususiy klinika esa xavfli belgini bir necha "
                            "soniyada ko'radi. Bugun sizning 5 ta mezoningiz bo'yicha birma-bir javob beramiz." % PRODUCT)
    label = "CHECKPOINT 2  ·  %s" % d.mentor.upper()
    box(s, 0.7, 0.7, 4.0, 0.46, fill=TEAL, radius=0.5)
    text(s, 0.7, 0.7, 4.0, 0.46, label, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.7, 1.6, 5.6, 1.2, PRODUCT, size=72, color=WHITE, bold=True, spacing=0.9)
    text(s, 0.7, 3.0, 5.3, 1.5, tagline, size=21, color="CFE3E6", spacing=1.15)
    text(s, 0.7, 4.85, 5.3, 0.6, [[("Qoida qaror qiladi. ", {"bold": True, "color": CORAL}), ("AI tushuntiradi.", {"bold": True, "color": WHITE})]], size=20)
    text(s, 0.7, 6.55, 6, 0.4, "%s  ·  Tibbiyot treki  ·  %s" % (TEAM, DATE), size=13, color=PALE)
    picture(s, image, 6.55, 1.5, 6.1)
    text(s, 6.45, 5.05, 6.3, 0.4, "Ishlayotgan mahsulot: chapda ona telefoni, o'ngda klinika paneli", size=12, color=PALE, align=PP_ALIGN.CENTER)


def map_slide(d, criteria):
    s = d.slide(WHITE, notes="Avval mahsulot uch qadamda, keyin har bir mezon uchun bitta slayd. Oxirida jonli demo.")
    text(s, 0.7, 0.55, 12, 0.4, "10 DAQIQALIK REJA  ·  %s" % d.mentor.upper(), size=13, color=TEAL, bold=True)
    text(s, 0.7, 0.9, 12, 0.8, "5 mezon, 5 ta aniq javob", size=34, bold=True)
    box(s, 0.7, 2.0, 4.6, 4.6, fill=LIGHT)
    text(s, 1.0, 2.2, 4, 0.4, "MAHSULOT QANDAY ISHLAYDI", size=12, color=MUTED, bold=True)
    steps = [("Ona 2 daqiqada javob beradi", "8–10 ta savol, o'zbek tilida, katta tugmalar"),
             ("Qoidalar xavfni aniqlaydi", "23 ta qoida: yashil / sariq / qizil"),
             ("Klinika 3 soniyada ko'radi", "Ogohlantirish, grafiklar, AI xulosasi")]
    for i, (t1, t2) in enumerate(steps):
        y = 2.8 + i * 1.25
        badge(s, 1.0, y, str(i + 1), d=0.6, fill=CORAL, size=18)
        text(s, 1.8, y - 0.04, 3.4, 0.4, t1, size=15, bold=True)
        text(s, 1.8, y + 0.36, 3.3, 0.6, t2, size=13, color=MUTED)
    for i, (name, proof) in enumerate(criteria):
        y = 2.0 + i * 0.94
        badge(s, 5.75, y + 0.06, str(i + 1), d=0.62)
        text(s, 6.6, y, 6.1, 0.4, name, size=18, bold=True)
        text(s, 6.6, y + 0.38, 6.1, 0.5, proof, size=13.5, color=MUTED)
    footer(s, d)


def footer(s, d, dark=False):
    text(s, 0.7, 6.95, 12, 0.3, "%s  ·  Checkpoint 2  ·  %s  ·  %s" % (PRODUCT, d.mentor, DATE), size=10.5, color=PALE if dark else MUTED)


def criterion(d, n, name, headline, notes, bg=WHITE):
    s = d.slide(bg, notes=notes)
    badge(s, 0.7, 0.55, str(n), d=0.78, size=26)
    text(s, 1.7, 0.5, 10, 0.4, "MEZON %d / 5  ·  %s" % (n, name.upper()), size=13, color=TEAL, bold=True)
    text(s, 1.7, 0.85, 10.9, 1.05, headline, size=29, bold=True, spacing=1.0)
    footer(s, d)
    return s


def cards(s, items, y=2.05, h=2.1, cols=None, fill=LIGHT, title_size=15, body_size=12):
    cols = cols or len(items)
    w = (11.95 - 0.25 * (cols - 1)) / cols
    for i, (title, body) in enumerate(items):
        x, yy = 0.7 + (i % cols) * (w + 0.25), y + (i // cols) * (h + 0.22)
        box(s, x, yy, w, h, fill=fill)
        text(s, x + 0.22, yy + 0.18, w - 0.44, 0.7, title, size=title_size, bold=True, color=TEAL_DARK, spacing=1.05)
        text(s, x + 0.22, yy + 0.18 + (0.72 if len(title) > 26 else 0.45), w - 0.44, h - 0.8, body, size=body_size, spacing=1.12)


def stats(s, items, y=2.05):
    w = (11.95 - 0.25 * (len(items) - 1)) / len(items)
    for i, (num, label, sub) in enumerate(items):
        x = 0.7 + i * (w + 0.25)
        box(s, x, y, w, 2.05, fill=TEAL if i == 0 else LIGHT)
        c = WHITE if i == 0 else INK
        size = 44 if len(num) <= 5 else 38 if len(num) <= 7 else 30 if len(num) <= 9 else 24
        text(s, x + 0.22, y + 0.12, w - 0.4, 0.8, num, size=size, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.22, y + 0.95, w - 0.4, 0.5, label, size=13, bold=True, color=c, spacing=1.0)
        text(s, x + 0.22, y + 1.47, w - 0.4, 0.55, sub, size=10.5, color="CFE3E6" if i == 0 else MUTED, spacing=1.0)


def table(s, head, rows, widths, y=2.0, row_h=0.62, size=12.5, strong_last=True):
    xs = [0.95]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    box(s, 0.7, y, 11.95, 0.5, fill=DARK, radius=0.15)
    for x, w, label in zip(xs, widths, head):
        text(s, x, y, w - 0.2, 0.5, label, size=12.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    for i, row in enumerate(rows):
        yy = y + 0.62 + i * (row_h + 0.04)
        if i % 2 == 0:
            box(s, 0.7, yy, 11.95, row_h, fill=LIGHT, radius=0.12)
        for j, (x, w, cell) in enumerate(zip(xs, widths, row)):
            last = j == len(row) - 1 and strong_last
            text(s, x, yy, w - 0.2, row_h, cell, size=size, bold=(j == 0 or last), color=TEAL_DARK if j == 0 else (INK if last else MUTED), anchor=MSO_ANCHOR.MIDDLE)


def closing(d, demo, nxt, image="patient.jpg"):
    s = d.slide(DARK, notes="Endi 2 daqiqalik jonli demo: telefonda so'rovnoma, qon ketishi belgisi, klinika panelida qizil ogohlantirish va AI xulosasi. Keyin savollaringizga javob beramiz.")
    text(s, 0.7, 0.55, 8, 0.4, "JONLI DEMO  ·  2 DAQIQA", size=13, color=CORAL, bold=True)
    text(s, 0.7, 0.9, 6.0, 1.1, "Gapirib bermaymiz, ko'rsatamiz", size=30, color=WHITE, bold=True, spacing=1.0)
    for i, (t1, t2) in enumerate(demo):
        y = 2.1 + i * 1.02
        badge(s, 0.7, y, str(i + 1), d=0.6, fill=CORAL, size=18)
        text(s, 1.55, y - 0.05, 4.9, 0.4, t1, size=17, bold=True, color=WHITE)
        text(s, 1.55, y + 0.33, 4.9, 0.6, t2, size=12.5, color=PALE)
    picture(s, image, 7.05, 0.8, 5.55)
    text(s, 6.95, 4.15, 5.75, 0.4, "KEYINGI QADAMLAR", size=13, color=CORAL, bold=True)
    bullets(s, 6.97, 4.62, 5.7, nxt, size=13.5, gap=0.45, color=WHITE)
    text(s, 0.7, 6.6, 12, 0.4, "%s  ·  %s" % (TEAM, REPO), size=13, color=PALE)


DEMO = [("Klinika: yangi bemor", "Ro'yxatga olish, QR kod. Ona telefonida skanerlaydi, o'rnatish shart emas"),
        ("Telefon: so'rovnoma", "«Qon ketishi» belgilanadi, izoh yoziladi. Qizil natija va qo'ng'iroq tugmasi"),
        ("Klinika: ≤ 3 soniya", "Qizil ogohlantirish izoh bilan; qabul qilish va onaga javob yozish"),
        ("AI xulosasi", "Gulnora: bosim 125/79 → 152/98, preeklampsiya xavfi, tavsiya etilgan choralar")]

# =====================================================================================
# 1. TEXNIK MENTOR
# =====================================================================================
d = Deck("Texnik mentor")
title_slide(d, "Homilador ayol uchun kunlik AI hamroh. Klinika uchun jonli kuzatuv paneli.")
map_slide(d, [("Texnik amalga oshirish", "Boshidan oxirigacha ishlaydi: ilova → qoidalar → AI → klinika paneli"),
              ("Kod tayyorligi", "13 avtomatik test, 18 API endpoint, bitta buyruq bilan ishga tushadi, GitHub'da"),
              ("Yechimning innovatsionligi", "Xavfni AI emas, qoida belgilaydi; chatdagi xavfli belgi klinikaga yetadi"),
              ("Jamoaning bilim darajasi", "Har bir chegara, prompt va cheklovni tushuntirib bera olamiz"),
              ("Texnologiyalar to'plami", "Python + SQLite + PWA + gpt-5-mini: yengil, arzon, provayderga bog'lanmagan")])

s = criterion(d, 1, "Texnik amalga oshirish", "Boshidan oxirigacha ishlaydigan tizim, maket emas",
              "Ona ilovasi javoblarni API'ga yuboradi. Avval qoidalar xavf darajasini hisoblaydi, keyin AI shu darajani ona tilida tushuntiradi. "
              "Natija bazaga yoziladi, klinika paneli uni 3 soniya ichida ko'rsatadi. AI yoki internet ishlamasa, shablon javoblar bilan davom etadi.")
flow = [("Ona ilovasi", "Mobil PWA\nuz / ru / en", SOFT, INK), ("API server", "18 ta endpoint\nJSON", SOFT, INK), ("Xavf qoidalari", "23 ta qoida\n13 ta test", TEAL, WHITE),
        ("AI qatlami", "gpt-5-mini\nstructured outputs", CORAL, WHITE), ("Klinika paneli", "jonli ogohlantirish\n≤ 3 soniya", SOFT, INK)]
for i, (t1, t2, fill, color) in enumerate(flow):
    x = 0.7 + i * 2.5
    box(s, x, 2.1, 2.1, 1.45, fill=fill)
    text(s, x + 0.12, 2.24, 1.86, 0.4, t1, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(s, x + 0.12, 2.68, 1.86, 0.8, t2.split("\n"), size=12, color=color if fill != SOFT else MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        box(s, x + 2.14, 2.72, 0.32, 0.24, fill=MUTED, shape=MSO_SHAPE.RIGHT_ARROW)
cards(s, [("3 ta AI funksiya jonli", "Onaga shaxsiy xabar, shifokor uchun 14 kunlik xulosa, chat va klinikaga eskalatsiya"),
          ("Xavfsizlik to'ri", "Qon ketishi, bosim ≥160/110, harakat yo'qligi qoida bilan ushlanadi, AI'ga bog'liq emas"),
          ("Oflayn ham ishlaydi", "AI yoki internet uzilsa, har bir funksiya qoidaga asoslangan shablonga o'tadi"),
          ("To'liq klinika oqimi", "Login, bemorni ro'yxatga olish, QR taklif, qidiruv va filtrlar, xabarlashuv")], y=4.0, h=2.6)

s = criterion(d, 2, "Kod tayyorligi", "Klon qiling, bitta buyruq bering: ishlaydi",
              "Kod GitHub'da, README bo'yicha bitta buyruq bilan ishga tushadi. Xavf qoidalari 13 ta avtomatik test bilan qoplangan. Kalitlar .env faylida, repozitoriyga tushmaydi.")
stats(s, [("13", "avtomatik test", "xavf qoidalari uchun, hammasi o'tadi"), ("18", "API endpoint", "har biri skript va brauzerda tekshirilgan"),
          ("~3 100", "qator kod", "Python, JavaScript, CSS"), ("1", "buyruq", "python server.py: baza va demo ma'lumot o'zi yaratiladi")])
text(s, 0.7, 4.2, 5.7, 0.4, "Repozitoriyda nima bor", size=16, bold=True, color=TEAL_DARK)
bullets(s, 0.72, 4.68, 5.7, ["README: o'rnatish va 3 daqiqalik demo ssenariysi", "requirements.txt va .env.example", "seed.py --reset: toza demo ma'lumotlar",
                             "CLAUDE.md: spetsifikatsiya va 12 bosqichli reja", "Modullar: triage / ai / db / server alohida"], gap=0.4)
text(s, 6.8, 4.2, 5.7, 0.4, "Xavfsizlik va ishonchlilik", size=16, bold=True, color=TEAL_DARK)
bullets(s, 6.82, 4.68, 5.8, ["Parollar: PBKDF2-SHA256, 200 000 iteratsiya", "Sessiya: HttpOnly cookie; panel API'lari sessiyasiz 401",
                             "Ona havolasi: taxmin qilib bo'lmaydigan maxfiy token", "API kalit faqat .env da, .gitignore ichida",
                             "AI chaqiruvi: 45 s timeout, xatoda shablon"], gap=0.4)
box(s, 8.75, 0.55, 3.9, 0.42, fill=DARK, radius=0.5)
text(s, 8.75, 0.55, 3.9, 0.42, REPO, size=11.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

s = criterion(d, 3, "Yechimning innovatsionligi", "Qoida qaror qiladi, AI tushuntiradi",
              "Tibbiyotda AI'ga qaror bermaymiz. Xavf darajasini test qilingan qoidalar belgilaydi, AI uni pasaytira olmaydi. Ona chatda xavfli belgini yozsa, "
              "klinika avtomatik ogohlantiriladi. Javob bermay qo'ygan bemor ham panelda ko'rinadi.")
table(s, ["", "Oddiy AI sog'liq chatboti", PRODUCT],
      [("Xavf darajasini kim belgilaydi", "Til modeli: xato qilishi mumkin", "JSST belgilari asosidagi 23 qoida, 13 test"),
       ("Chatda xavfli belgi aytilsa", "Javob beradi, xolos", "Klinikaga avtomatik ogohlantirish ketadi"),
       ("Bemor javob bermay qo'ysa", "Hech kim bilmaydi", "3 kundan keyin «javobsiz» belgisi"),
       ("Internet yoki AI uzilsa", "Xizmat to'xtaydi", "Shablon rejimida ishlashda davom etadi"),
       ("Vaqt bo'yicha tahlil", "Bir martalik savol-javob", "3 kunlik ketma-ketlik, haftalik vazn, 14 kunlik AI xulosa"),
       ("Shaffoflik", "Qora quti", "«Klinik asos» sahifasi: har bir qoida, sharti va manbasi")], [3.3, 4.1, 4.4])

s = criterion(d, 4, "Jamoaning bilim darajasi", "Har bir qarorni tushuntirib bera olamiz, cheklovlarni ham",
              "Bu slayd siz berishingiz mumkin bo'lgan savollarga javoblarimiz. Chegaralar qayerdan olingani, AI qanday nazorat qilinishi va bugungi cheklovlarni ochiq aytamiz.", bg=LIGHT)
cards(s, [("Nega xavfni AI emas, qoida belgilaydi?", "Tibbiyotda xato narxi yuqori. Qoidani test qilish va shifokor bilan tekshirish mumkin. AI faqat tushuntiradi."),
          ("Chegaralar qayerdan olingan?", "JSST xavf belgilari; bosim ≥140/90 sariq, ≥160/110 qizil (ISSHP/ACOG); harakat 24-haftadan; to'lg'oq 37-haftagacha qizil."),
          ("AI javobi qanday nazorat qilinadi?", "Structured outputs (JSON schema). Prompt: tashxis yo'q, dori yo'q. Rad javobi va uzilish tekshiriladi, xatoda shablon."),
          ("AI'ga qanday ma'lumot ketadi?", "Ism (familiyasiz), yosh, hafta, javoblar. Telefon, familiya, manzil yuborilmaydi. Keyingi qadam: ismni ham olib tashlash."),
          ("Ma'lumotlar qanday himoyalangan?", "PBKDF2 parollar, HttpOnly sessiya, maxfiy taklif tokeni, kalitlar .env da. Demo ma'lumotlari sintetik."),
          ("Bugungi cheklovlar nimalar?", "Chegaralar mahalliy akusher-ginekolog bilan tasdiqlanmagan. gpt-5-mini o'zbekchasi ba'zan noaniq. Tibbiy qurilma emas.")],
      y=2.05, h=2.2, cols=3, fill=WHITE, title_size=14, body_size=11.5)

s = criterion(d, 5, "Texnologiyalar to'plami", "Yengil, arzon va provayderga bog'lanmagan",
              "Har bir texnologiya aniq sabab bilan tanlangan: tez ishga tushishi, arzonligi va bitta AI provayderga bog'lanib qolmaslik. AI qatlami OpenAI va Anthropic'ni bir xil interfeys orqali qo'llaydi.")
table(s, ["Qatlam", "Texnologiya", "Nega aynan shu"],
      [("Mobil ilova", "HTML, CSS, JavaScript (PWA)", "Har qanday telefonda brauzerdan ochiladi, do'kondan o'rnatish shart emas"),
       ("Backend", "Python 3.9, standart kutubxona", "Freymvorksiz: bitta buyruq, har qanday noutbukda ishlaydi"),
       ("Ma'lumotlar bazasi", "SQLite", "Sozlash kerak emas; sxema PostgreSQL ga oson ko'chadi"),
       ("Sun'iy intellekt", "OpenAI gpt-5-mini, rasmiy SDK", "Tez (3–13 s) va arzon; Anthropic Claude zaxira provayder sifatida tayyor"),
       ("Xavf dvigateli", "Sof Python qoidalari + unittest", "Shaffof, tekshiriladigan, AI'dan mustaqil"),
       ("Grafiklar va QR", "Inline SVG; segno (QR)", "Tashqi JS kutubxonasiz; 140/90 chegarasi chizig'i bilan"),
       ("Xavfsizlik", "PBKDF2-SHA256, HttpOnly cookie", "Standart kutubxona vositalari")], [2.5, 3.9, 5.4], row_h=0.56, strong_last=False)
closing(d, DEMO, ["Telegram orqali kunlik eslatmalar va shifokorga ogohlantirish", "Ona uchun telefon raqami + SMS-kod bilan kirish",
                  "AI'ga ismsiz (psevdonim) ma'lumot; lokal ochiq model varianti", "API va autentifikatsiya uchun avtotestlar"])
d.save("NigohAI_CP2_1_Texnik_mentor.pptx")

# =====================================================================================
# 2. BIZNES MENTOR
# =====================================================================================
d = Deck("Biznes mentor")
title_slide(d, "Xususiy klinikalar uchun homiladorlikni masofadan kuzatish xizmati. Ona uchun bepul, klinika uchun obuna.")
map_slide(d, [("G'oya va muammoning dolzarbligi", "Ayollar ko'rikka keladi, lekin tashriflar orasida klinika hech narsani ko'rmaydi"),
              ("Biznes modeli va monetizatsiya", "B2B2C: klinika obuna to'laydi, ona bepul foydalanadi"),
              ("Raqobatdagi ustunlik", "Yagona mahsulot: ona ilovasi + klinika paneli + xavf qoidalari, o'zbek tilida"),
              ("Moliyaviy barqarorlik", "Past xarajat: AI bemor boshiga oyiga ≈ 1–2 ming so'm, do'kon komissiyasi yo'q"),
              ("O'sish imkoniyati", "8 700 xususiy klinika; homiladorlikdan bola salomatligigacha kengayish")])

s = criterion(d, 1, "G'oya va muammoning dolzarbligi", "Muammo tibbiyotga kirishda emas, tashriflar orasidagi bo'shliqda",
              "O'zbekistonda ayollarning deyarli hammasi homiladorlik ko'rigidan o'tadi. Lekin ko'riklar orasida 2–4 hafta o'tadi, preeklampsiya va homila harakati "
              "kamayishi esa kunlar ichida rivojlanadi. Shu oraliqda klinika hech narsani ko'rmaydi.")
stats(s, [("98,8%", "kamida 1 marta ko'rikdan o'tgan", "yaqinda tug'gan ayollar"), ("89,5%", "4 va undan ko'p ko'rik", "tizimga kirish muammo emas"),
          ("53,5%", "8 va undan ko'p ko'rik", "JSST tavsiyasiga yarmi yetadi"), ("2–4 hafta", "ko'riklar orasidagi ko'r hudud", "xavfli belgilar kunlarda rivojlanadi")])
text(s, 0.7, 4.17, 12, 0.3, "Manba: O'zbekiston MICS 2021–2022 (IPUMS Microdata Center)", size=10.5, color=MUTED)
cards(s, [("Ona uchun", "«Bu normalmi? Qachon va kimga murojaat qilaman?» degan savolga tez javob yo'q. Qishloqdagi ayol klinikadan uzoqda."),
          ("Klinika uchun", "Bemor uyda nima bo'layotganini bilmaydi. Xavfli holatni faqat keyingi tashrifda yoki tez yordamda ko'radi."),
          ("Bizning g'oya", "Har kuni 2 daqiqalik so'rovnoma. Xavfli belgi klinikaga 3 soniyada yetadi. AI tushuntiradi, shifokor qaror qiladi.")], y=4.55, h=2.1)

s = criterion(d, 2, "Biznes modeli va monetizatsiya", "B2B2C: ona bepul foydalanadi, klinika obuna to'laydi",
              "Biz bemor ma'lumotini sotmaymiz. Klinika parvarishni boshqarish vositasi uchun to'laydi: 9 oy davomida bemorni ushlab qoladi, xavfni erta ko'radi va xizmat sifatini ko'rsatadi.")
cards(s, [("1. SaaS obuna (asosiy)", "Klinika har bir faol homilador uchun oylik to'lov to'laydi: panel, ogohlantirishlar, AI xulosalar, xabarlashuv, QR bilan ulash."),
          ("2. Homiladorlik paketlari", "Klinika + Nigoh AI: 9 oylik raqamli va klinik paket. Klinika paketni qimmatroq sotadi, biz platforma to'lovini olamiz."),
          ("3. Kelajak: B2B va davlat", "Sug'urta va ish beruvchilar uchun homiladorlik dasturlari; davlat uchun onalik salomatligi monitoringi shartnomalari.")], y=2.05, h=2.35)
text(s, 0.7, 4.7, 6, 0.4, "Klinika nega to'laydi", size=16, bold=True, color=TEAL_DARK)
bullets(s, 0.72, 5.15, 5.9, ["Bemorni 9 oy ushlab qoladi, bir martalik qabul emas", "Xavfli holatni erta ko'radi: obro' va xavfsizlik",
                             "Shifokor vaqtini tejaydi: kimga e'tibor kerakligi tayyor"], gap=0.42)
box(s, 6.9, 4.7, 5.75, 1.9, fill=DARK)
text(s, 7.15, 4.85, 5.3, 1.6, [[("Biz shaxsiy tibbiy ma'lumotni sotmaymiz.", {"bold": True, "color": WHITE})],
                               [("Daromad parvarishni muvofiqlashtirishdan keladi, bemor ma'lumotidan emas.", {"color": PALE})]], size=15)

s = criterion(d, 3, "Raqobatdagi ustunlik", "Infratuzilma bilan raqobat emas: tashriflar orasidagi qatlam",
              "DMED sog'liqni saqlashni raqamlashtiradi, biz uni almashtirmaymiz, to'ldiramiz. Xorijiy homiladorlik trekerlari esa faqat ona uchun: klinika bilan bog'lanmagan va o'zbek tilida emas.")
table(s, ["", "DMED", "Xorijiy homiladorlik trekerlari", PRODUCT],
      [("Asosiy vazifa", "Umumiy tibbiy infratuzilma", "Ona uchun kalendar va maqolalar", "Tashriflar orasida kuzatuv"),
       ("Kunlik xavf so'rovnomasi", "Yo'q", "Qisman, klinikaga yetmaydi", "Bor: 23 ta qoida, 3 daraja"),
       ("Klinikaga jonli ogohlantirish", "Yo'q", "Yo'q", "Bor: ≤ 3 soniya"),
       ("Shifokor uchun AI xulosa", "Umumiy AI vositalari", "Yo'q", "Bor: 14 kunlik trend va choralar"),
       ("O'zbek tili, o'rnatishsiz", "O'zbek tilida", "Yo'q", "Bor: QR orqali brauzerda"),
       ("Mijoz", "Davlat va muassasalar", "Iste'molchi (B2C)", "Xususiy klinikalar (B2B2C)")], [2.9, 2.7, 3.2, 3.0], size=12)

s = criterion(d, 4, "Moliyaviy barqarorlik", "Past o'zgaruvchan xarajat, obunaga asoslangan daromad",
              "Xarajat tuzilmasi yengil: AI arzon model, ilova do'konlari komissiyasi yo'q, bitta server yetarli. Narx hozircha gipoteza, uni pilotda 1–2 klinika bilan tekshiramiz.")
stats(s, [("≈ 1–2 ming", "so'm: AI xarajati", "bemor boshiga oyiga, gpt-5-mini (taxminiy hisob)"), ("0%", "do'kon komissiyasi", "PWA: App Store va Google Play kerak emas"),
          ("1 server", "yuzlab bemor uchun", "Python + SQLite, keyin PostgreSQL"), ("4 kishi", "jamoa", "dasturchi, loyiha rahbari, biznes, tibbiyot")])
box(s, 0.7, 4.2, 11.95, 2.45, fill=LIGHT)
text(s, 1.0, 4.35, 11, 0.4, "BIRLIK IQTISODIYOTI  ·  GIPOTEZA, PILOTDA TEKSHIRILADI", size=12, bold=True, color=MUTED)
text(s, 1.0, 4.8, 11.4, 0.5, "Oylik daromad = klinikalar soni × faol homiladorlar × bemor boshiga oylik narx", size=17, bold=True, color=TEAL_DARK)
text(s, 1.0, 5.4, 11.4, 1.2, ["Misol: 1 klinika × 200 homilador × 20 000 so'm = 4 mln so'm / oy.  25 klinika = 100 mln so'm / oy.",
                              "AI va server xarajati daromadning 10% idan kam bo'lishi kutiladi. Asosiy xarajat: jamoa va klinikalarni ulash (sotuv)."], size=13.5)

s = criterion(d, 5, "O'sish imkoniyati", "O'sayotgan xususiy tibbiyot bozori va kengayadigan mahsulot",
              "Xususiy tibbiyot bozori tez o'smoqda. Biz homiladorlikdan boshlaymiz, lekin xuddi shu halqa tug'ruqdan keyingi davr va bola salomatligiga ham cho'ziladi: o'sha klinikalar, o'sha oilalar.")
stats(s, [("19,9 trln", "so'm: tibbiy xizmatlar, 2025", "+14,2% yiliga (Milliy statistika qo'mitasi)"), ("~30%", "xususiy sektor ulushi", "2025-yil (President.uz)"),
          ("3 200 → 8 700", "xususiy klinikalar", "2016 → 2026 oxiri, kutilmoqda (President.uz)"), ("2030", "davlat maqsadi", "xususiy klinikalarni davlat buyurtmasiga jalb qilish")])
text(s, 0.7, 4.2, 12, 0.4, "Kengayish yo'li", size=16, bold=True, color=TEAL_DARK)
path = ["Pilot: 1–2 klinika, Xorazm", "Viloyat klinikalari", "Respublika + DMED integratsiyasi", "Tug'ruqdan keyin va chaqaloq", "Bola salomatligi, Markaziy Osiyo"]
for i, step in enumerate(path):
    x = 0.7 + i * 2.43
    box(s, x, 4.75, 2.2, 1.35, fill=TEAL if i == 0 else LIGHT)
    text(s, x + 0.15, 4.85, 1.9, 0.35, "%d-bosqich" % (i + 1), size=11, bold=True, color="CFE3E6" if i == 0 else MUTED)
    text(s, x + 0.15, 5.2, 1.9, 0.85, step, size=13.5, bold=True, color=WHITE if i == 0 else INK)
    if i < 4:
        box(s, x + 2.22, 5.3, 0.2, 0.22, fill=MUTED, shape=MSO_SHAPE.RIGHT_ARROW)
text(s, 0.7, 6.3, 12, 0.4, "Pilotda o'lchanadi: ogohlantirishga javob vaqti, so'rovnoma to'ldirish ulushi, klinikaning to'lashga tayyorligi.", size=13, color=MUTED)
closing(d, DEMO, ["1–2 xususiy klinika bilan pilot va narxni tekshirish", "Telegram eslatmalar: so'rovnoma to'ldirish ulushini oshirish",
                  "Klinika uchun tahlil sahifasi: javob vaqti, faollik", "DMED va klinika tizimlari bilan integratsiya"])
d.save("NigohAI_CP2_2_Biznes_mentor.pptx")

# =====================================================================================
# 3. SOHA MENTORI
# =====================================================================================
d = Deck("Soha mentori")
title_slide(d, "Tashriflar orasida homilador ayolni kuzatish: xavf belgisi klinikaga bir necha soniyada yetadi.")
map_slide(d, [("Muammoning dolzarbligi", "Xavfli holatlar kunlarda rivojlanadi, ko'riklar orasi esa 2–4 hafta"),
              ("Sohani chuqur tushunish", "23 ta qoida: har birining sharti va manbasi shifokor o'qiy oladigan ko'rinishda"),
              ("Amaliy qo'llanish", "QR bilan ulash, 2 daqiqalik so'rovnoma, shifokorga tayyor xulosa"),
              ("Qonunchilikka muvofiqlik", "Tashxis qo'ymaydi; ma'lumot minimal; shaxsiy ma'lumotlar qonuni hisobga olingan"),
              ("Ta'sir ko'lami", "Har bir xususiy klinika va uzoq tumandagi har bir homilador ayol")])

s = criterion(d, 1, "Muammoning dolzarbligi", "Xavfli belgilar kunlarda rivojlanadi, klinika esa haftalab ko'rmaydi",
              "JSST ma'lumotiga ko'ra qon ketishi va gipertenziv buzilishlar onalar o'limining yetakchi sabablaridan. Ikkalasi ham erta belgilar beradi, "
              "lekin ayol ularni klinikaga yetkazmasa, hech kim bilmaydi.")
cards(s, [("Preeklampsiya", "Bosim, bosh og'rig'i, ko'rish buzilishi, shish bir necha kunda kuchayadi. Uy sharoitidagi bosim ko'rsatkichi va belgilar erta signal beradi."),
          ("Homila harakati kamayishi", "Ona sezadi, lekin «keyingi ko'rikkacha kutaman» deb o'ylaydi. Ko'rsatma esa kutmaslikni aytadi."),
          ("Qon ketishi va suv ketishi", "Daqiqalar muhim. Ayolga aniq yo'riq kerak: hoziroq qo'ng'iroq qiling, 103."),
          ("Ruhiy holat va vitaminlar", "Past kayfiyat, xavotir, temir preparatini ichmaslik: bir kunlik emas, ketma-ket kunlarda ko'rinadi.")], y=2.05, h=2.35)
box(s, 0.7, 4.75, 11.95, 1.85, fill=DARK)
text(s, 1.0, 4.9, 11.4, 1.6, [[("98,8% ayol ko'rikdan o'tadi, 53,5% 8 va undan ko'p marta (MICS 2021–2022).", {"bold": True, "color": WHITE})],
                              [("Demak muammo tizimga kirishda emas. Muammo: tashriflar orasida nima bo'layotganini hech kim ko'rmaydi.", {"color": PALE})]], size=15.5)

s = criterion(d, 2, "Sohani chuqur tushunish", "Har bir qoida, uning sharti va manbasi ochiq ko'rsatilgan",
              "Klinik asos sahifasida 23 ta qoidaning hammasi bor: qaysi shartda ishlaydi va qaysi manbaga tayanadi. To'rtta qoida bizning evristikamiz, "
              "buni ham ochiq yozganmiz. Chegaralarni mahalliy akusher-ginekologlar bilan tasdiqlash keyingi qadam.")
picture(s, "rules.jpg", 6.75, 2.1, 5.85)
bullets(s, 0.72, 2.1, 5.7, ["Qizil: qon ketishi, suv ketishi, bosim ≥160/110, harakat yo'q, kuchli qorin og'rig'i, nafas qisishi",
                            "Preeklampsiya: 20-haftadan bosim ≥140/90 + bitta belgi yoki ikkita belgi birga",
                            "Haftaga bog'liq mantiq: harakat 24-haftadan; to'lg'oq 37-haftagacha qizil, keyin sariq",
                            "Uydagi o'lchov xatolari: 60–260 / 30–160 dan tashqari bosim e'tiborga olinmaydi",
                            "Trendlar: 3 kun past kayfiyat, 3 kun vitamin ichilmagan, haftada 2 kg dan ortiq vazn",
                            "Manbalar: JSST xavf belgilari, ISSHP/ACOG, RCOG; 4 ta qoida evristika deb belgilangan"], size=13, gap=0.72)

s = criterion(d, 3, "Amaliy qo'llanish", "Klinikaning bugungi ish jarayoniga 5 daqiqada qo'shiladi",
              "Hamshira bemorni ro'yxatga oladi va QR kod beradi. Ona hech narsa o'rnatmaydi. Shifokor ertalab kimga e'tibor kerakligini ko'radi, "
              "bemor sahifasida esa tayyor xulosa, grafiklar va ayolning o'z so'zlari bor.")
picture(s, "patient.jpg", 6.75, 2.1, 5.85)
flow = [("Ro'yxatga olish + QR", "Ism, hafta, xavf omillari. Ona kamerasi bilan skanerlaydi."),
        ("Kunlik 2 daqiqa", "Katta tugmalar, belgilar rasm bilan, o'zbek / rus tili. Savodi past ayol uchun ham qulay."),
        ("Ogohlantirish ≤ 3 soniya", "Qizil va sariq holatlar, ayolning izohi bilan. «Qabul qilindi» tugmasi."),
        ("Shifokor sahifasi", "Oxirgi bosim, vazn, 14 kunlik grafik, AI xulosa va tavsiya etilgan choralar."),
        ("Javobsiz bemor", "3 kun javob bermasa panelda alohida ko'rinadi: qo'ng'iroq qilish uchun signal.")]
for i, (t1, t2) in enumerate(flow):
    y = 2.05 + i * 0.92
    badge(s, 0.7, y, str(i + 1), d=0.56, fill=CORAL, size=16)
    text(s, 1.45, y - 0.05, 5.0, 0.35, t1, size=15, bold=True)
    text(s, 1.45, y + 0.28, 5.0, 0.6, t2, size=12, color=MUTED)

s = criterion(d, 4, "Qonunchilikka muvofiqlik", "Tashxis qo'ymaydi, ma'lumotni minimallashtiradi, qaror shifokorda",
              "Tizim qaror qabul qilishga yordam beradi, tashxis qo'ymaydi va dori tavsiya qilmaydi. Shaxsga doir ma'lumotlar to'g'risidagi qonun talablarini "
              "hisobga oldik. Pilotdan oldin yurist va Sog'liqni saqlash vazirligi talablari bo'yicha tekshiruv o'tkazamiz.", bg=LIGHT)
cards(s, [("Tibbiy qaror shifokorda", "Ilova tashxis qo'ymaydi, dori yoki doza tavsiya qilmaydi. Onaga faqat: «klinikaga qo'ng'iroq qiling / 103». Har bir AI xulosasida ogohlantirish bor."),
          ("Shaxsga doir ma'lumotlar (O'RQ-547)", "Baza lokal: klinika yoki O'zbekiston hududidagi serverda joylashtiriladi. Kirish login bilan, parollar xeshlangan, ona havolasi maxfiy token."),
          ("AI'ga minimal ma'lumot", "Tashqi AI'ga familiya, telefon, manzil yuborilmaydi: faqat ism, yosh, hafta va javoblar. Keyingi qadam: ismsiz yuborish yoki lokal model."),
          ("Rozilik", "Keyingi qadam: ro'yxatga olishda ayolning yozma roziligi va ma'lumotlardan foydalanish shartlari. Hozir demo ma'lumotlari sintetik."),
          ("Tibbiy buyum maqomi", "Bugun: qaror qabul qilishni qo'llab-quvvatlash vositasi, tibbiy qurilma emas. Maqomni yurist bilan aniqlashtiramiz."),
          ("Klinik tasdiqlash", "Chegaralar xalqaro ko'rsatmalardan. Pilotdan oldin milliy protokollar va mahalliy akusher-ginekologlar bilan tasdiqlanadi.")],
      y=2.05, h=2.2, cols=3, fill=WHITE, title_size=14, body_size=11.5)

s = criterion(d, 5, "Ta'sir ko'lami", "Bitta klinikadan butun onalik parvarishi zanjirigacha",
              "Ta'sir uch darajada: ona xavfli belgini erta biladi, klinika kimga e'tibor kerakligini ko'radi, tizim esa og'ir holatlarni kamaytiradi. "
              "Xuddi shu yondashuv tug'ruqdan keyingi davr va chaqaloq parvarishiga ham o'tadi.")
cards(s, [("Ona", "Har kuni 2 daqiqa. Xavfli belgida aniq yo'riq. Savoliga o'z tilida javob. Uzoq tumanda ham klinika bilan aloqada."),
          ("Shifokor va klinika", "Kimga bugun e'tibor kerakligi tayyor. 14 kunlik trend va AI xulosa: 20 soniyada qaror. Javobsiz bemorlar ko'rinadi."),
          ("Sog'liqni saqlash tizimi", "Erta aniqlash: kamroq og'ir asorat va tez yordam chaqiruvi. Ma'lumot: tashriflar orasidagi uzluksiz kuzatuv.")], y=2.05, h=2.2)
stats(s, [("8 700", "xususiy klinika", "2026 oxiriga kutilmoqda (President.uz)"), ("9 oy +", "har bir ayol bilan aloqa", "homiladorlik, keyin tug'ruqdan keyingi davr"),
          ("3 til", "o'zbek, rus, ingliz", "o'rnatishsiz, har qanday telefonda"), ("Pilot", "1–2 klinika", "o'lchov: javob vaqti, faollik, o'tkazib yuborilgan holatlar")], y=4.55)
closing(d, DEMO, ["Chegaralarni milliy protokol va ginekologlar bilan tasdiqlash", "Homila harakatini sanash (kick counter) va tashrif eslatmalari",
                  "Yozma rozilik va ma'lumotlardan foydalanish shartlari", "Tug'ruqdan keyingi davr va chaqaloq uchun so'rovnomalar"])
d.save("NigohAI_CP2_3_Soha_mentori.pptx")
